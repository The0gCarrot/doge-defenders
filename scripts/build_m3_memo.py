"""Build the M3 memo from the notebook's measured results and saved figures."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "docs" / "m3_figures"
OUTS = [
    ROOT / "docs" / "BUA3336_M3_Doge_Defenders.docx",
    Path(r"c:\Users\gabri\Downloads\BUA3336_M3_Doge_Defenders.docx"),
]

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLACK = RGBColor(0x22, 0x22, 0x22)


def set_run(run, size=10.5, bold=False, italic=False, color=BLACK):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_text(paragraph, text, size=10.5, bold=False, italic=False, color=BLACK):
    run = paragraph.add_run(text)
    set_run(run, size, bold, italic, color)
    return run


def p(doc, text="", size=10, bold=False, italic=False, space_before=1, space_after=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if text:
        add_text(paragraph, text, size, bold, italic)
    return paragraph


def heading(doc, text):
    return p(doc, text, size=11, bold=True, space_before=4, space_after=1)


def caption(doc, text):
    return p(doc, text, size=9, italic=True, space_before=1, space_after=4)


def shade(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): hex_color,
    })
    tcPr.append(shd)


def table(doc, headers, rows):
    grid = doc.add_table(rows=1 + len(rows), cols=len(headers))
    grid.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = grid.rows[0].cells[i]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        add_text(paragraph, header, size=8, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "1F4E79")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = grid.rows[r + 1].cells[c]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            add_text(paragraph, value, size=8)
            if r % 2 == 1:
                shade(cell, "F4F7FB")
    return grid


def picture(doc, name, width):
    doc.add_picture(str(FIG / name), width=Inches(width))
    doc.paragraphs[-1].paragraph_format.space_before = Pt(1)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)


def picture_pair(doc, left, right, width=3.35):
    grid = doc.add_table(rows=1, cols=2)
    for cell, name in zip(grid.rows[0].cells, (left, right)):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run()
        run.add_picture(str(FIG / name), width=Inches(width))
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "nil")
            borders.append(element)
        tcPr.append(borders)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    p(doc, "BUA 3336 · Big Data & Data Mining for Business", size=11, bold=True, space_before=0)
    p(doc, "M3 — Exploratory Analysis Report", size=14, bold=True, space_before=0, space_after=2)
    p(doc, "Team: Doge Defenders. Data Steward: Daniel Burke. Modeling Lead: Gabriel McLaughlin. AI Auditor: Augusta Domingo. Business Translator: Mason Ocampo. QA Lead: no separate named seat. The four members review one another's checks.")
    p(doc, "Author of this memo: Gabriel McLaughlin. Date: October 1, 2026.")
    p(doc, "This memo looks at the file before any model is trusted. The file has 525 rows. One row is one coin on one day. Comparisons with the 90-day result use the 510 rows where that result is already known.")

    heading(doc, "Part 0 — Revised target")
    p(doc, "The target is the answer we are trying to predict. We are predicting whether the coin's price is higher about 90 days later, for each daily report (one coin on one day), using only information available at midnight UTC on that report date. UTC is the world's standard clock, so \"today\" is the same moment for every coin.")
    p(doc, "The cutoff rule is the exact line between the two answers. We set it in the M2 report and kept it. The result is up (1) when the later price is higher than the price on the report date. Otherwise it is not up (0). If the future price does not exist yet, we do not call it up. We ignore trading fees. The 15 rows dated September 26, 2026 stay out of Parts 2 through 5, because 90 days have not passed.")
    p(doc, "A labeled row is one that already has that answer. Of the 510 labeled rows, 406 finished up (79.6%). That is the majority, the result that happened more often. 104 did not finish up (20.4%). That is the minority, the smaller group. The course asks for at least 60 in the smaller group. We have 104.")

    heading(doc, "Part 1 — The notebook")
    p(doc, "Notebook link: https://colab.research.google.com/github/The0gCarrot/doge-defenders/blob/m3-exploratory/notebooks/M3_exploratory_analysis.ipynb")
    p(doc, "A notebook is a list of steps a computer runs from top to bottom. Colab is the free website that runs it. Use Runtime, then Restart and run all. That clones the m3-exploratory branch and reads data/processed/daily_reports.csv. Every number in this memo is printed by that notebook. Above each block of code, a note says what the block does and why we ran it.")

    heading(doc, "Part 2 — Univariate and bivariate exploration")
    p(doc, "2.1  One column at a time. Univariate means looking at one column by itself. The mean is the ordinary average. The median is the middle value after the rows are lined up from smallest to largest. Skew describes a lopsided shape. A right skew means a few very large numbers pull the average above the middle row. When the mean and the median are far apart, the average is describing a rare coin, not a typical day. These six numbers are the ones a manager would be tempted to quote. Both the mean and the median come from the 510 labeled rows.", space_after=3)
    table(
        doc,
        ["Variable", "Mean", "Median", "Shape", "What the mean hides"],
        [
            ["Price", "$4,581", "$2.95", "Right skew 3.53", "BTC and ETH are 99.7% of summed prices. A typical coin costs about $3."],
            ["Dollar volume", "$3.59 billion", "$295 million", "Right skew 3.45", "The average day is about twelve times the middle day."],
            ["Market value", "$114.3 billion", "$3.86 billion", "Right skew 3.38", "Bitcoin is 78.0% of summed market value. The average describes Bitcoin."],
            ["30-day change", "−14.3%", "−18.1%", "Right skew 2.94", "The middle coin was down about 18%. Rallies up to +123% lift the average."],
            ["Headline count", "1.52", "0", "Right skew 1.77", "54.3% of rows have no headline. The average implies a story the middle day lacks."],
            ["24-hour return", "−0.76%", "−0.70%", "Skew 0.06", "The average does not mislead. A typical day was a small loss."],
        ],
    )
    picture(doc, "01_market_cap_distribution.png", 3.9)
    caption(doc, "Figure 1. Half of the coin-days are worth under about $3.9 billion. The $114 billion average sits far to the right because Bitcoin pulls it there. When you describe a typical coin, quote the median, not the mean. The horizontal axis is a log scale: each step multiplies the dollars by 10, so a coin worth $1 billion and a coin worth $100 billion can sit on the same chart.")

    p(doc, "2.2  Each predictor against the result. Bivariate means two columns at once. A predictor is a fact we already knew on the report date. The result is the answer 90 days later. For a number, we compare the median on the 406 up rows with the median on the 104 not-up rows. For a label, such as the coin's category, we report the share that finished up. A quartile is one quarter of the rows after they are sorted. The \"weakest\" quarter is the 128 rows that had already fallen the most over the prior month.", space_before=4, space_after=3)
    table(
        doc,
        ["Predictor", "How we compared", "What we see"],
        [
            ["30-day change", "Median by class; share up by quarter", "Median −20.4% when up, −5.4% when not. Weakest 128 rows: 125 up (97.7%). Strongest 128: 67 up (52.3%)."],
            ["Vs 30-day high", "Median by class", "Up rows were further below the high (median −23.6%) than not-up rows (−15.1%)."],
            ["Category", "Share that finished up", "DeFi 102/102. Layer 1: 222/272 (81.6%). Meme: 43/68 (63.2%). Smart-contract platform: 39/68 (57.4%). DeFi here is only Chainlink, Aave, and Uniswap."],
            ["Headline opinion", "Share that finished up", "No headline: 229/277 (82.7%). Any headline: 177/233 (76.0%). Bullish 62/78 (79.5%). Bearish 26/34 (76.5%). Neutral 89/121 (73.6%)."],
            ["24-hour return", "Median by class", "−0.67% when up, −0.76% when not. The straight-line link with the result is −0.002, which is no link."],
        ],
    )
    picture_pair(doc, "02_momentum_vs_target.png", "03_up_rate_by_date.png", 3.2)
    p(doc, "A few labels in that table need a plain meaning. Layer 1 is a base blockchain, such as Bitcoin. DeFi, short for decentralized finance, means lending and trading tools that are not a bank. Bullish means the headlines sound like good news for the price. Bearish means they sound like bad news. Neutral means the words are mixed, or the two news sources do not point the same way.", space_before=2)
    caption(doc, "Figure 2 (left). Coins that had just had their best month were higher 90 days later 67 of 128 times. Coins that had fallen the hardest were higher 125 of 128 times. The chart rounds those shares to 52% and 98%. Figure 3 (right). June 2 was the low point: only 46.7% of the 15 coins were higher 90 days later. Before June 10, 166 of 255 rows were higher (65.1%). From June 10 on, 240 of 255 were higher (94.1%). The dashed line is June 10.")

    heading(doc, "Part 3 — Correlation and similarity")
    p(doc, "3.1  The heatmap. A correlation is a score from −1 to +1. It asks whether two number columns rise and fall together in a straight line. Near +1, they move together. Near −1, one rises while the other falls. Near 0, there is no straight-line link. A heatmap is that grid drawn in color. Red means the columns move together. Blue means they move apart. The map uses the 510 labeled rows and nine number columns. It leaves out the 90-day result, the model scores, and the practice portfolio. Those are the answer, or they were built from the answer. Using them would be like seeing the test key before the test.")
    p(doc, "Strongest relationship 1: price and market value, correlation +0.990. Market value is price times how many coins exist, so this line is mostly the definition of market value. It is not a discovery about trading.")
    p(doc, "Strongest relationship 2: dollars traded and market value, correlation +0.880. Price and dollars traded are next, at +0.838. The 30-day price change and the distance from the 30-day high move together at +0.632. Those two facts are partly the same fact written twice.")
    p(doc, "What the heatmap does not tell you. It does not say that one column causes the other. It only notices straight lines, so a relationship that rises and then falls can look weak. It leaves out category, the exchange, and the headline text, because those are names and sentences, not numbers. It compares columns. It does not say whether two coin-days are alike. And the darkest square is not the useful one for our question. The 30-day price change is almost unrelated to market value (about −0.01 on the map), but it is the strongest number-column link to the 90-day result (−0.409).")
    picture(doc, "04_correlation_heatmap.png", 3.7)
    caption(doc, "Figure 4. Price, dollars traded, and market value are one cluster. A number is printed only where the correlation is at least 0.50 in size, ignoring the sign. A red square is not a reason to buy.")

    p(doc, "3.2  Similarity. Correlation compares columns. A trading question often compares rows: is this coin-day like that coin-day? One of our rows holds 9 number fields, 3 name fields (category, exchange, and headline opinion), a date, a list of up to 30 past prices, and headline text. That mix is why the measure matters.")
    p(doc, "Euclidean distance is the ordinary straight-line gap. You subtract every number, square each gap, add them, and take the square root. On the raw file that fails, because market value is in the billions and a percent is a small number. The billions swallow the distance. On May 24, 2026, the raw distance from Bitcoin to Pepe is about 1.54 trillion, and 99.96% of the squared gap is market value alone. Squared means the gap is multiplied by itself before it is added, so the huge gap dominates even more. Bitcoin to Ethereum is about 1.28 trillion. Raw Euclidean treats Ethereum as almost as different from Bitcoin as Pepe is.")
    p(doc, "The fitting measure for a mixed row is Gower distance. On that same day, each number gap is divided by that column's own range across the 15 coins, so a percent and a market value are put on the same 0-to-1 scale. A name either matches (0) or it does not (1). Every field counts once. Zero would mean the rows match. One would mean they share nothing. Gower scores Bitcoin–Ethereum at 0.218, Bitcoin–Pepe at 0.598, and Dogecoin–Pepe at 0.285. Dogecoin and Pepe land closer because they share the meme label. Jaccard distance only counts which yes-or-no labels overlap, so it would throw away the dollar amounts. Cosine similarity compares the direction of two text lists, so it fits headlines, not prices.")

    heading(doc, "Part 4 — Three business findings")
    p(doc, "Each finding is a sentence a manager could act on. The evidence is the chart or table behind it, and the row count is how many daily reports support it.")
    p(doc, "4.1  Do not treat a strong recent month as a reason to buy. Figure 2 is the evidence. The 128 rows with the weakest 30-day change finished higher 125 times. The 128 rows with the strongest 30-day change finished higher 67 times. That is 256 rows inside the 510 labeled rows. A manager would stop using \"it just rallied\" as a buy signal in this book, and would read Part 5 before turning the chart into a rule.")
    p(doc, "4.2  Do not give Dogecoin and Pepe one meme budget. Figure 5 is the evidence. Dogecoin was higher on 10 of its 34 labeled days (29.4%). Pepe was higher on 33 of 34 (97.1%). The meme average, 43 of 68 days (63.2%), describes neither coin. A manager would size the two coins separately. The category word is not the position.")
    picture(doc, "05_up_rate_by_token.png", 3.5)
    caption(doc, "Figure 5. Seven coins were higher on all 34 labeled days: Bitcoin, Ethereum, Solana, XRP, Chainlink, Uniswap, and Aave. Dogecoin, in rust, was the weakest name in the file. Pepe, in teal, was close to that perfect group. Both carry the meme label.")
    p(doc, "4.3  Do not let a headline overturn the price view in this file. The opinion table in section 2.2 is the evidence. Days with no headline finished higher 229 of 277 times (82.7%). Days with at least one headline finished higher 177 of 233 times (76.0%). Bullish days (79.5%) and bearish days (76.5%) are close. All 510 labeled rows are behind this. A manager would keep any news adjustment small. In this window, the tone of the headline does not separate the coins that rose from the coins that did not.")

    heading(doc, "Part 5 — The surprise")
    p(doc, "What surprised us. We expected the coins that had just risen to be the ones higher 90 days later. The strongest link to the result is the 30-day price change, and it points the other way. The correlation is −0.409. A negative correlation means that as the recent month looks better, the chance of being higher in 90 days looks worse.")
    p(doc, "The exciting explanation is that the market punished people who chased a hot month and rewarded the coins that had already fallen. The boring explanation is the calendar. Later June start dates were deeper in a selloff. The 30-day change correlates −0.360 with the report date, so later days tend to show a worse recent month. Those same dates almost all finished up. The result correlates +0.404 with the report date. A drop followed by a recovery makes \"was down recently\" and \"was up 90 days later\" show up together, even if the recent drop is not itself a buy signal.")
    p(doc, "We believe the boring explanation is the main reason the full-file number looks so strong. We checked by splitting the 510 rows at June 10. Before that date, the weakest quarter finished higher 60 of 64 times and the strongest quarter finished higher 28 of 64 times, so a gap is still there early. On and after June 10, the weakest quarter was 64 of 64 and the strongest quarter was 63 of 64. The late window does not show the pattern, because almost every coin finished up. We would not hand a manager Figure 2 as a standing rule.")

    heading(doc, "Required plots")
    p(doc, "Figure 1 is the distribution, the shape of one column: the average coin is not the typical coin. Figure 2 is the relationship: the strongest recent months were the group least often higher 90 days later. Figure 3 is the time comparison: late-June start dates almost all finished higher. Figure 4 is the heatmap: price, dollars traded, and market value repeat one another. Figure 5 is our own chart: Dogecoin and Pepe share a label and not an outcome.")

    heading(doc, "Part 6 — AI Collaboration Log")
    p(doc, "Copied forward from the Week 4 decision log in the M2 report. These are mistakes we caught and checked. They are not a new log written for this memo.")
    p(doc, "Exchanges and categories. We asked an assistant to fill in the main exchange and the coin's category. It picked the trading pair with the biggest reported volume and the first tag on CoinGecko's list. Odd exchanges looked like Bitcoin's home market, and the names of investment funds looked like business types. We opened the raw lists and checked them by hand. The file now uses Binance, KuCoin, and Coinbase International, and four types: Layer 1, DeFi, Meme, and Smart Contract Platform.")
    p(doc, "News source. We asked for headlines from the last 10 days on CoinGecko. The assistant used a status-update link that no longer exists. Every coin returned \"page not found.\" We repeated the calls ourselves and saw the same error. We did not make up headlines. The team switched to GDELT, a news search of stories from around the world, and Yahoo Finance. If one source sounds positive and the other sounds negative, the opinion is neutral and the score changes by 0.")
    p(doc, "Cleaning and prices. A first draft wanted empty columns dropped, extreme prices cut off, and the whole file rescaled before any rows were set aside for testing. Rescaling means shrinking very large numbers so an $80,000 price does not overwhelm a 2% change. The test split means some rows are held back to check the model, instead of letting it study every row. We kept every row, marked holes as −999 for unknown numbers, NULL for unknown words, and N/A for portfolio cells before the start day, and rescaled only inside the model after that split. A saved file from September 17 priced Bitcoin near $76,608. A live check the same week was about $84,000, so today's row uses the live price. The rule \"score divided by 10, as a percent of the cash still left\" was our rule, not a default from the assistant.")

    for path in OUTS:
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            doc.save(path)
            print("Wrote", path)
        except PermissionError:
            alt = path.with_name(path.stem + "_plain" + path.suffix)
            doc.save(alt)
            print("Locked; wrote", alt)


if __name__ == "__main__":
    main()
